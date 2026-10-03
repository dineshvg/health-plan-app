"""Build the tracker workbook (.xlsx) from your PRIVATE tracker settings.

    pip install openpyxl
    cp tracker/settings.example.json tracker/settings.json   # git-ignored; fill in your numbers
    python tracker/build_tracker.py                           # writes tracker/tracker.xlsx
    python tracker/build_tracker.py path/to/settings.json out.xlsx

Settings hold weights, goal, milestones and habits, so they never go into the public app/.

Upload the .xlsx to Google Drive and open it with Google Sheets (File > Save as Google Sheets).
The app reads cells by position, so keep the layout described in tracker/LAYOUT.md.
"""
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
if len(sys.argv) > 1:
    src = Path(sys.argv[1])
else:
    src = ROOT / "tracker" / "settings.json"
    if not src.exists():
        src = ROOT / "tracker" / "settings.example.json"
        print("tracker/settings.json not found, using the example settings")
T = json.loads(src.read_text())
OUT = sys.argv[2] if len(sys.argv) > 2 else str(ROOT / "tracker" / "tracker.xlsx")
START = dt.date.fromisoformat(T["startDate"])
DAYS, WEEKS = 364, 52
HABIT = T["habit"]
RULES = (T["dailyRules"] + [""] * 6)[:6]  # exactly 6 Y/N columns (J-O)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2E5E4E")
NOTE = Font(italic=True, color="666666")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
iso = dt.date.fromisoformat

wb = Workbook()


def header(ws, cols, row=1, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = (widths or {}).get(c, max(11, min(len(c) + 2, 22)))
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def yn(ws, rng):
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Y"'], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N"'], fill=RED))


last7 = lambda col: f"=IFERROR(AVERAGEIFS('Daily Log'!{col}:{col},'Daily Log'!A:A,\">\"&TODAY()-7,'Daily Log'!A:A,\"<=\"&TODAY()),\"-\")"

# ---------- Dashboard (rows 2-13 and 16-20 are read by the app) ----------
d = wb.active
d.title = "Dashboard"
for col, wdt in zip("ABCD", (34, 16, 16, 18)):
    d.column_dimensions[col].width = wdt
d["A1"] = T.get("name", "Health plan")
d["A1"].font = Font(bold=True, size=14)
rows = [
    ("Start date", START),                                                                   # B2
    ("Start weight (kg)", T["startWeight"]),                                                 # B3
    ("Goal weight (kg)", T["goalWeight"]),                                                   # B4
    ("Latest 7-day average (kg)", "=IFERROR(LOOKUP(9.99E+307,'Daily Log'!D:D,'Daily Log'!E:E),B3)"),
    ("Lost so far (kg)", "=B3-B5"),
    ("Still to go (kg)", "=B5-B4"),
    ("Progress to goal", "=IF(B3=B4,0,B6/(B3-B4))"),                                         # B8
    ("Workouts done", "=COUNTIF(Workouts!F:F,\"Y\")"),
    ("Workouts planned up to today", "=COUNTIFS(Workouts!B:B,\"<=\"&TODAY())"),
    (HABIT["weekly"] + ", last 7 days", last7("G")),
    ("Avg sleep (h), last 7 days", last7("H")),
    (HABIT["freeDays"], "=COUNTIF('Daily Log'!G:G,0)"),                                       # B13
]
for i, (k, v) in enumerate(rows, 2):
    d.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = d.cell(row=i, column=2, value=v)
    c.number_format = "DD.MM.YYYY" if i == 2 else "0%" if i == 8 else "0.0"
d.cell(row=15, column=1, value="Milestones").font = Font(bold=True, size=12)
for j, h in enumerate(["Milestone", "Date", "Target (kg)", "Actual avg, 7 days to date"], 1):
    c = d.cell(row=16, column=j, value=h)
    c.font, c.fill = HEAD, FILL
for k, (name, date, tgt) in enumerate(T["milestones"][:4]):
    rr = 17 + k
    d.cell(row=rr, column=1, value=name)
    d.cell(row=rr, column=2, value=iso(date)).number_format = "DD.MM.YYYY"
    d.cell(row=rr, column=3, value=tgt)
    d.cell(row=rr, column=4, value=f"=IF(B{rr}>TODAY(),\"\",IFERROR(LOOKUP(B{rr},'Daily Log'!A:A,'Daily Log'!E:E),\"\"))").number_format = "0.0"
d.cell(row=22, column=1, value="Log one row a day in 'Daily Log', tick sessions in 'Workouts', review on Sunday. Formula cells update themselves.").font = NOTE

# ---------- Daily Log (columns A-S are read and written by the app) ----------
s = wb.create_sheet("Daily Log")
cols = ["Date", "Day", "Week", "Weight (kg)", "7-day avg (kg)", "Steps", HABIT["label"], "Sleep (h)",
        "Sleep quality 1-5", *RULES, T.get("drinks", "Alcohol drinks"), "Water (L)", "Rules hit %", "Notes / mood"]
header(s, cols, widths={"Notes / mood": 40, "Rules hit %": 10})
for i in range(DAYS):
    r = i + 2
    day = START + dt.timedelta(days=i)
    s.cell(row=r, column=1, value=day).number_format = "DD.MM.YYYY"
    s.cell(row=r, column=2, value=day.strftime("%a"))
    s.cell(row=r, column=3, value=i // 7 + 1)
    s.cell(row=r, column=5, value=f'=IF(D{r}="","",AVERAGE(D{max(2, r - 6)}:D{r}))').number_format = "0.0"
    s.cell(row=r, column=18, value=f'=IF(COUNTA(J{r}:O{r})=0,"",COUNTIF(J{r}:O{r},"Y")/COUNTA(J{r}:O{r}))').number_format = "0%"
yn(s, f"J2:O{DAYS + 1}")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
s.add_data_validation(dv)
dv.add(f"I2:I{DAYS + 1}")
s.conditional_formatting.add(f"A2:S{DAYS + 1}", FormulaRule(formula=["$A2=TODAY()"], fill=PatternFill("solid", fgColor="FFF2CC")))

# ---------- Workouts (A-E planned, F-K written by the app) ----------
w = wb.create_sheet("Workouts")
header(w, ["Week", "Planned date", "Phase", "Session", "Focus", "Done (Y/N)", "Duration (min)",
           "Hardest RPE", "Key lift + weight", "Rounds / laps", "Notes"],
       widths={"Focus": 34, "Key lift + weight": 24, "Notes": 36, "Planned date": 13})
r = 2
n_weeks = T.get("weeks", 12)
for wk in range(1, n_weeks + 1):
    phase = f"{(wk - 1) * 3 // n_weeks + 1}"
    for sess in T["sessions"]:
        w.cell(row=r, column=1, value=wk)
        w.cell(row=r, column=2, value=START + dt.timedelta(days=(wk - 1) * 7 + sess["dayOffset"])).number_format = "DD.MM.YYYY"
        w.cell(row=r, column=3, value="Phase " + phase)
        w.cell(row=r, column=4, value=sess["code"])
        w.cell(row=r, column=5, value=sess["focus"])
        r += 1
yn(w, "F2:F500")

# ---------- Weekly Review (row = week + 1; F, N, O written by the app) ----------
v = wb.create_sheet("Weekly Review")
header(v, ["Week", "Mon", "Sun (review day)", "Avg weight (kg)", "Change vs last week", "Waist (cm)",
           "Avg sleep (h)", HABIT["weekly"], "Avg steps", "Workouts done", "Rules hit %",
           T.get("drinks", "Alcohol drinks"), "3 flat weeks?", "Win of the week", "Fix for next week"],
       widths={"Win of the week": 34, "Fix for next week": 34})
DL = "'Daily Log'"
for wk in range(1, WEEKS + 1):
    r = wk + 1
    avg = lambda col: f'=IFERROR(AVERAGEIFS({DL}!{col}:{col},{DL}!C:C,A{r}),"")'
    v.cell(row=r, column=1, value=wk)
    v.cell(row=r, column=2, value=START + dt.timedelta(days=(wk - 1) * 7)).number_format = "DD.MM"
    v.cell(row=r, column=3, value=START + dt.timedelta(days=(wk - 1) * 7 + 6)).number_format = "DD.MM"
    v.cell(row=r, column=4, value=avg("D")).number_format = "0.0"
    v.cell(row=r, column=5, value=f'=IF(OR(D{r}="",D{r - 1}=""),"",D{r}-D{r - 1})' if wk > 1 else "").number_format = "+0.0;-0.0;0.0"
    v.cell(row=r, column=7, value=avg("H")).number_format = "0.0"
    v.cell(row=r, column=8, value=avg("G")).number_format = "0.0"
    v.cell(row=r, column=9, value=avg("F")).number_format = "#,##0"
    v.cell(row=r, column=10, value=f'=COUNTIFS(Workouts!A:A,A{r},Workouts!F:F,"Y")')
    v.cell(row=r, column=11, value=avg("R")).number_format = "0%"
    v.cell(row=r, column=12, value=f'=SUMIFS({DL}!P:P,{DL}!C:C,A{r})')
    if wk >= 4:
        v.cell(row=r, column=13, value=f'=IF(OR(D{r}="",D{r - 3}=""),"",IF(MIN(D{r - 2}:D{r})>=D{r - 3}-0.1,"ADJUST","ok"))')
v.conditional_formatting.add(f"M2:M{WEEKS + 1}", CellIsRule(operator="equal", formula=['"ADJUST"'], fill=RED))

# ---------- Tests & Photos ----------
t = wb.create_sheet("Tests & Photos")
header(t, ["Test", "Week 1", f"Week {n_weeks}", "Change"], widths={"Test": 40})
for i, tst in enumerate(T.get("tests", []), 2):
    t.cell(row=i, column=1, value=tst)
    t.cell(row=i, column=4, value=f'=IF(OR(B{i}="",C{i}=""),"",C{i}-B{i})')

wb.save(OUT)
print("saved", OUT)
